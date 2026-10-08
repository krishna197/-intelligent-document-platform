import io
import os

from docx import Document
from fastapi import APIRouter, File, HTTPException, UploadFile
from pypdf import PdfReader

router = APIRouter()

ALLOWED_TYPES = {".pdf", ".docx", ".txt"}
MAX_BYTES = 10 * 1024 * 1024  # 10 MB


def extract_pdf(data: bytes) -> tuple[str, int]:
    reader = PdfReader(io.BytesIO(data))
    pages = [(page.extract_text() or "").strip() for page in reader.pages]
    return "\n\n".join(p for p in pages if p), len(pages)


def extract_docx(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def extract_txt(data: bytes) -> str:
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        return data.decode("latin-1")


@router.post("/upload")
async def upload_document(file: UploadFile = File(...)):
    """Upload a PDF, DOCX or TXT file and return its extracted text."""
    ext = os.path.splitext(file.filename or "")[1].lower()
    if ext not in ALLOWED_TYPES:
        raise HTTPException(status_code=400, detail="Unsupported file type. Use PDF, DOCX or TXT.")

    data = await file.read()
    if len(data) > MAX_BYTES:
        raise HTTPException(status_code=413, detail="File is larger than 10 MB.")

    pages = None
    try:
        if ext == ".pdf":
            text, pages = extract_pdf(data)
        elif ext == ".docx":
            text = extract_docx(data)
        else:
            text = extract_txt(data)
    except Exception as e:
        raise HTTPException(status_code=422, detail=f"Could not read this file: {e}")

    text = text.strip()
    if not text:
        raise HTTPException(
            status_code=422,
            detail="No text found. If this is a scanned PDF, it needs OCR, which is not supported yet.",
        )

    metadata = {"size": len(data), "type": ext, "words": len(text.split())}
    if pages is not None:
        metadata["pages"] = pages

    return {"filename": file.filename, "text": text, "metadata": metadata, "status": "success"}